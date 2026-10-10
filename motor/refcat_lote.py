# -*- coding: utf-8 -*-
"""
Referencia catastral masiva: Excel del cliente -> Excel con el código predial de cada fila,
más un KML con los casos por revisar ya ubicados (punto estimado y predios candidatos).
"""
from __future__ import annotations

import os
from collections import Counter
from xml.sax.saxutils import escape

import geopandas as gpd
import openpyxl
from openpyxl.styles import Font, PatternFill
from shapely.geometry import Point

from .base import cargar_base
from .excel import leer_excel
from .refcat import Geocodificador

COLORES = {"ALTA": "C6EFCE", "MEDIA": "FFEB9C", "BAJA": "FFC7CE", "SIN_RESULTADO": "D9D9D9"}


def resolver_filas(filas, geo: Geocodificador):
    for f in filas:
        r = geo.buscar(f.get("NOMENCLATU", ""), cedula=f.get("NO_DE_IDEN"))
        f["_REF"] = r
    return filas


def escribir_excel(filas, ruta):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "REFERENCIA_CATASTRAL"
    enc = ["HOJA", "FILA", "NODO", "NOMBRE", "CEDULA", "DIRECCION", "BARRIO", "REF_CATASTRAL", "METODO", "CONFIANZA",
           "REVISAR", "CANDIDATOS", "NOTA", "REF_ORIGINAL"]
    ws.append(enc)
    for c in ws[1]:
        c.font = Font(bold=True)
    for f in filas:
        r = f["_REF"]
        cands = "; ".join(f"{c} ({m})" for c, m in r.candidatos if c != r.codigo)
        ws.append([f["_HOJA"], f["_FILA"], f.get("NODO", ""), f.get("NOMBRE_BEN", ""), f.get("NO_DE_IDEN", ""),
                   f.get("_DIRECCION_COMPLETA") or f.get("NOMENCLATU", ""), f.get("BARRIO_URB", ""), r.codigo, r.metodo,
                   r.confianza, "" if r.confianza == "ALTA" else "SI", cands, r.nota, f.get("CODIGO", "")])
        fill = PatternFill("solid", fgColor=COLORES.get(r.confianza, "FFFFFF"))
        for col in (8, 10, 11):
            ws.cell(ws.max_row, col).fill = fill
        ws.cell(ws.max_row, 8).number_format = "@"
    for col, ancho in zip("ABCDEFGHIJKLMN", (16, 6, 8, 34, 13, 24, 16, 33, 22, 14, 9, 60, 60, 33)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"
    res = wb.create_sheet("RESUMEN")
    cuenta = Counter(f["_REF"].confianza for f in filas)
    res.append(["CONFIANZA", "FILAS", "QUÉ HACER"])
    for k, que in (("ALTA", "Lista: placa del catastro o dirección del R1 confirmada"),
                   ("MEDIA", "Revisar en el KML: hay un candidato probable"),
                   ("BAJA", "Revisar en el KML: ubicación aproximada, elegir entre los candidatos"),
                   ("SIN_RESULTADO", "Dirección incompleta o vía que no existe en la base")):
        res.append([k, cuenta.get(k, 0), que])
    res.column_dimensions["A"].width, res.column_dimensions["C"].width = 16, 70
    wb.save(ruta)
    return cuenta


def escribir_kml(filas, geo: Geocodificador, ruta):
    """Casos por revisar: pin en el punto estimado y polígonos de los candidatos (para Google Earth)."""
    base = geo.base
    terr = base.terrenos.set_index("codigo")
    partes = ['<?xml version="1.0" encoding="UTF-8"?>', '<kml xmlns="http://www.opengis.net/kml/2.2"><Document>',
              "<name>Referencias por revisar</name>",
              '<Style id="c1"><LineStyle><color>ff00a5ff</color><width>3</width></LineStyle><PolyStyle><color>5500a5ff</color></PolyStyle></Style>',
              '<Style id="c2"><LineStyle><color>ff00ffff</color><width>2</width></LineStyle><PolyStyle><color>3300ffff</color></PolyStyle></Style>']
    n = 0
    for f in filas:
        r = f["_REF"]
        if r.confianza == "ALTA" or (not r.punto and not r.candidatos):
            continue
        n += 1
        nombre = escape(f"{f.get('NODO', '')} {f.get('NOMENCLATU', '')}".strip())
        desc = escape(f"{f.get('NOMBRE_BEN', '')} | propuesto: {r.codigo} [{r.metodo}/{r.confianza}] {r.nota}")
        partes.append(f"<Folder><name>{nombre}</name>")
        if r.punto:
            p = gpd.GeoSeries([Point(r.punto)], crs=base.crs).to_crs(4326).iloc[0]
            partes.append(f"<Placemark><name>{nombre}</name><description>{desc}</description>"
                          f"<Point><coordinates>{p.x:.7f},{p.y:.7f},0</coordinates></Point></Placemark>")
        for i, (cod, motivo) in enumerate(r.candidatos[:6]):
            t = geo.terreno_de(cod)
            if not t or t not in terr.index:
                continue
            g = gpd.GeoSeries([terr.geometry[t]], crs=base.crs).to_crs(4326).iloc[0]
            for pol in (g.geoms if hasattr(g, "geoms") else [g]):
                coords = " ".join(f"{x:.7f},{y:.7f},0" for x, y in pol.exterior.coords)
                partes.append(f"<Placemark><name>{cod}</name><description>{escape(motivo)}</description>"
                              f"<styleUrl>#{'c1' if cod == r.codigo else 'c2'}</styleUrl><Polygon><outerBoundaryIs>"
                              f"<LinearRing><coordinates>{coords}</coordinates></LinearRing></outerBoundaryIs></Polygon></Placemark>")
        partes.append("</Folder>")
    partes.append("</Document></kml>")
    with open(ruta, "w", encoding="utf-8") as fh:
        fh.write("\n".join(partes))
    return n


def procesar(ruta_excel, carpeta_salida=None, municipio="malambo"):
    base = cargar_base(municipio)
    geo = Geocodificador(base)
    filas = resolver_filas(leer_excel(ruta_excel), geo)
    carpeta_salida = carpeta_salida or os.path.dirname(os.path.abspath(ruta_excel))
    os.makedirs(carpeta_salida, exist_ok=True)
    nombre = os.path.splitext(os.path.basename(ruta_excel))[0]
    xlsx = os.path.join(carpeta_salida, f"{nombre}_REF_CATASTRAL.xlsx")
    kml = os.path.join(carpeta_salida, f"{nombre}_POR_REVISAR.kml")
    cuenta = escribir_excel(filas, xlsx)
    n_kml = escribir_kml(filas, geo, kml)
    return {"filas": len(filas), "confianza": dict(cuenta), "excel": xlsx, "kml": kml, "en_kml": n_kml}
