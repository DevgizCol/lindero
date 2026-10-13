# -*- coding: utf-8 -*-
"""
Flujo completo: Excel + plano del cliente -> capas, colindantes, PDF por predio y ZIP.

Etapas (cada una deja su resultado en el informe):
  1. excel        lee los Excel y los lleva a los campos de LEVANTAMIENTO_01
  2. referencia   código predial de las filas que no lo traen (placa / R1 / interpolación)
  3. poligonos    arma los polígonos del DXF (o lee una capa ya dibujada) y les pone el NODO
  4. ubicacion    detecta el sistema de coordenadas y ajusta sobre lc_terreno
  5. capas        LEVANTAMIENTO_01, COTAS y COORDENADAS con colindantes
  6. atlas        un PDF por NODO con la plantilla de QGIS, firma incluida, y el ZIP
"""
from __future__ import annotations

import datetime
import json
import os
import subprocess
import sys
import zipfile
from dataclasses import asdict, dataclass, field

import geopandas as gpd
import shapely

from . import capas as capas_mod
from .base import RAIZ, cargar_base
from .colindantes import VIA
from .excel import leer_excel
from .poligonos import Dibujo, desde_poligonos, extraer, norm_nodo, extraer_num_nodo
from .refcat import Geocodificador
from .ubicacion import ubicar

CAPAS_MAPA = ["COORDENADAS", "COTAS", "LEVANTAMIENTO_01", "08433_MALAMBO — cc_nomeclatura_vial", "lc_terreno"]


@dataclass
class Trabajo:
    carpeta: str
    excel: list = field(default_factory=list)
    dxf: str | None = None
    puntos: str | None = None          # cartera de campo (txt/csv) para etiquetar por código
    capa_poligonos: str | None = None  # alternativa al DXF: capa de polígonos con campo NODO
    crs_poligonos: str | None = None   # CRS de esa capa si no lo trae
    municipio: str = "malambo"
    firma: str | None = None           # None = la de la plantilla; "" = sin firma
    firmante: str | None = None
    fecha: str | None = None
    nombre: str = "ENTREGA"
    vista_previa: bool = False
    nodos_previa: int = 2
    hacer_atlas: bool = True


def _informe_vacio(t: Trabajo):
    return {"trabajo": t.nombre, "fecha": datetime.datetime.now().isoformat(timespec="seconds"), "etapas": {},
            "avisos": [], "lotes": [], "requiere_revision": False, "archivos": {}}


def _wkt_salida(cfg, base) -> str:
    """CRS de las capas de salida, escrito como lo espera el proyecto de QGIS (ESRI:103599 en Malambo)."""
    import pyproj
    crs = pyproj.CRS.from_user_input(cfg["crs_salida"]) if cfg.get("crs_salida") else base.crs
    return crs.to_wkt("WKT1_ESRI")


def ejecutar(t: Trabajo, registrar=print) -> dict:
    t.carpeta = os.path.abspath(t.carpeta)
    os.makedirs(t.carpeta, exist_ok=True)
    inf = _informe_vacio(t)
    base = cargar_base(t.municipio)
    cfg = base.cfg

    # 1. Excel
    filas = []
    for ruta in t.excel:
        filas += leer_excel(ruta)
    por_nodo = {}
    for f in filas:
        k = norm_nodo(f.get("NODO"))
        if not k:
            continue
        if k in por_nodo:
            inf["avisos"].append(f"NODO {f.get('NODO')} repetido en el Excel; se usa la primera fila.")
            continue
        por_nodo[k] = f
    inf["etapas"]["excel"] = {"filas": len(filas), "con_nodo": len(por_nodo)}
    registrar(f"Excel: {len(filas)} filas, {len(por_nodo)} con NODO")
    if not por_nodo:
        inf["avisos"].append("El Excel no trae la columna NODO (o viene vacía): no se pueden relacionar filas y polígonos.")

    # 2. Referencia catastral por dirección para las filas sin código
    geo = Geocodificador(base)
    ref = {"CLIENTE": 0, "ALTA": 0, "MEDIA": 0, "BAJA": 0, "SIN_RESULTADO": 0}
    for f in por_nodo.values():
        cod = "".join(c for c in str(f.get("CODIGO") or "") if c.isdigit())
        if len(cod) == 30 and geo.terreno_de(cod):
            f["_cod"], f["_cod_origen"], f["_cod_conf"] = cod, "CLIENTE", "ALTA"
            ref["CLIENTE"] += 1
            continue
        r = geo.buscar(f.get("NOMENCLATU", ""), cedula=f.get("NO_DE_IDEN"))
        f["_cod"], f["_cod_origen"], f["_cod_conf"], f["_cod_nota"] = r.codigo, r.metodo, r.confianza, r.nota
        f["_cod_cands"] = [c for c, _ in r.candidatos]
        ref[r.confianza] += 1
    inf["etapas"]["referencia"] = ref
    registrar(f"Referencia catastral: {ref}")

    # 3. Polígonos
    if t.capa_poligonos:
        g = gpd.read_file(t.capa_poligonos, engine="pyogrio")
        if g.crs is None and t.crs_poligonos:
            g = g.set_crs(t.crs_poligonos)
        campo = next((c for c in g.columns if c.upper() == "NODO"), None)
        dib = desde_poligonos(g.geometry, g[campo] if campo else [""] * len(g))
        crs_dibujo = str(g.crs) if g.crs is not None else None
    elif t.dxf:
        dib = extraer(t.dxf, [f.get("NODO") for f in por_nodo.values()], t.puntos)
        crs_dibujo = None
    else:
        dib = Dibujo(avisos=["No se entregó plano (DXF) ni capa de polígonos."])
        crs_dibujo = None
    inf["avisos"] += dib.avisos
    pares = []          # (nodo normalizado, polígono)
    vistos = set()
    polys_sin_emparejar = []
    for p, n in zip(dib.poligonos, dib.nodos):
        k = norm_nodo(n)
        if k and k in por_nodo and k not in vistos:
            pares.append((k, p))
            vistos.add(k)
        else:
            polys_sin_emparejar.append((p, n))

    # Rescate inteligente difuso: si hay polígonos sin emparejar y filas del Excel pendientes,
    # emparejar por número cardinal ('T1' con 'Lote 1', 'T-02' con '2', etc.)
    filas_pendientes = {k: f for k, f in por_nodo.items() if k not in vistos}
    if polys_sin_emparejar and filas_pendientes:
        mapa_num_excel = {}
        for k_ex, f_ex in filas_pendientes.items():
            num = extraer_num_nodo(f_ex.get("NODO", ""))
            if num:
                mapa_num_excel.setdefault(num, []).append(k_ex)

        for p, n in polys_sin_emparejar:
            num_poly = extraer_num_nodo(n)
            if num_poly and num_poly in mapa_num_excel:
                candidatos = mapa_num_excel[num_poly]
                if len(candidatos) == 1 and candidatos[0] not in vistos:
                    k_elegido = candidatos[0]
                    pares.append((k_elegido, p))
                    vistos.add(k_elegido)
                    inf["avisos"].append(
                        f"Emparejamiento inteligente: Polígono '{n}' asociado con éxito a fila '{por_nodo[k_elegido].get('NODO')}' por correspondencia de número #{num_poly}."
                    )

    sin_fila = [n for _, n in polys_sin_emparejar if not n or norm_nodo(n) not in vistos]
    sin_poly = [f.get("NODO") for k, f in por_nodo.items() if k not in vistos]
    if sin_fila:
        inf["avisos"].append(f"Polígonos cuyo NODO no está en el Excel: {', '.join(map(str, [x for x in sin_fila if x][:15]))}")
    if sin_poly:
        inf["avisos"].append(f"Filas del Excel sin polígono en el plano: {', '.join(map(str, sin_poly[:15]))}")
    inf["etapas"]["poligonos"] = {"en_plano": len(dib.poligonos), "con_nodo": sum(1 for n in dib.nodos if n),
                                  "emparejados": len(pares), "sin_poligono": sin_poly, "sin_fila": sin_fila}
    registrar(f"Polígonos: {len(dib.poligonos)} en el plano, {len(pares)} emparejados con el Excel")
    if not pares:
        inf["requiere_revision"] = True
        inf["avisos"].append("No hay ningún polígono emparejado con el Excel: verifica que los nombres de NODO coincidan o lleven la misma numeración.")
        return _cerrar(t, inf)

    # 4. Ubicación
    polys = [p for _, p in pares]
    cods = [por_nodo[k]["_cod"] if por_nodo[k].get("_cod_conf") == "ALTA" else "" for k, _ in pares]
    cods_terr = [geo.terreno_de(c) if c else "" for c in cods]
    cand = ([crs_dibujo] if crs_dibujo else []) + list(cfg.get("crs_origen_candidatos", []))
    ub = ubicar(polys, base, cods_terr, crs_candidatos=cand)
    inf["avisos"] += ub.avisos
    inf["etapas"]["ubicacion"] = {"metodo": ub.metodo, "crs_plano": ub.crs_origen, "dx": round(ub.dx, 2),
                                  "dy": round(ub.dy, 2), "giro": round(ub.giro, 2),
                                  "fuera_antes": round(ub.fuera_antes, 4), "fuera_despues": round(ub.fuera_despues, 4)}
    registrar(f"Ubicación: {ub.metodo} ({ub.crs_origen}), ajuste dx={ub.dx:.2f} dy={ub.dy:.2f} giro={ub.giro:.2f}°, "
              f"fuera de predio {100 * ub.fuera_antes:.1f}% -> {100 * ub.fuera_despues:.1f}%")
    if ub.metodo == "SIN_UBICAR":
        inf["requiere_revision"] = True
        return _cerrar(t, inf)
    if ub.metodo != "COORDENADAS_DEL_PLANO":
        inf["requiere_revision"] = True

    # código definitivo de cada lote: el del cliente/dirección si es confiable; si no, el predio que queda debajo
    lotes_info = []
    for (k, _), geom, debajo, calce in zip(pares, ub.poligonos, ub.codigo_debajo, ub.calce):
        f = por_nodo[k]
        cod, origen, conf = f.get("_cod", ""), f.get("_cod_origen", ""), f.get("_cod_conf", "")
        nota = ""
        if conf != "ALTA":
            if debajo and calce >= 0.6 and ub.metodo == "COORDENADAS_DEL_PLANO":
                coincide = cod and geo.terreno_de(cod) == debajo
                if not coincide:
                    cod = debajo
                origen = "TRASLAPE" + ("+DIRECCION" if coincide else "")
                conf = "ALTA" if (coincide or calce >= 0.85) else "MEDIA"
                nota = f"predio bajo el lote ({100 * calce:.0f}% del lote dentro)"
            elif not cod and debajo:
                cod, origen, conf = debajo, "TRASLAPE", "BAJA"
        elif debajo and geo.terreno_de(cod) != debajo:
            nota = f"el lote cae sobre {debajo} ({100 * calce:.0f}%), distinto al código que trae"
        info = dict(f)
        info.update(NODO=f.get("NODO"), geom=geom, CODIGO=cod)
        lotes_info.append(info)
        inf["lotes"].append({"nodo": f.get("NODO"), "codigo": cod, "codigo_origen": origen, "codigo_confianza": conf,
                             "codigo_nota": nota or f.get("_cod_nota", ""), "predio_debajo": debajo,
                             "calce": round(calce, 3), "area": round(geom.area, 2)})

    # 5. Capas y colindantes
    ctx = base.contexto()
    medidor = capas_mod.Medidor(base.crs, cfg.get("medidas", "elipsoidal"))
    lev, cotas, coord, detalle = capas_mod.construir(lotes_info, ctx, cfg["municipio"], medidor=medidor)
    carp_shp = os.path.join(t.carpeta, "01.SHAPEFILE")
    rutas = capas_mod.escribir_capas(carp_shp, lev, cotas, coord, _wkt_salida(cfg, base))
    inf["archivos"]["capas"] = rutas
    conf_lados = {"ALTA": 0, "MEDIA": 0, "BAJA": 0}
    for lote_inf, (lote, lados) in zip(inf["lotes"], detalle):
        lote_inf["lados"] = [{"tramo": f"{l.indice} AL {l.indice % len(lados) + 1}", "distancia": round(l.longitud, 2),
                              "colindante": l.colindante, "tipo": l.tipo, "lindero": l.lindero, "cardinal": l.cardinal,
                              "confianza": l.confianza, "fraccion": l.fraccion, "separacion": l.separacion,
                              "otros": l.secundarios} for l in lados if l.longitud >= 0.02]
        for l in lote_inf["lados"]:
            conf_lados[l["confianza"]] += 1
        lote_inf["frente_a_via"] = any(l.tipo == VIA for l in lados)
        lote_inf["revisar"] = (lote_inf["codigo_confianza"] != "ALTA" or lote_inf["calce"] < 0.85
                               or any(l["confianza"] == "BAJA" for l in lote_inf["lados"])
                               or "distinto al código" in lote_inf["codigo_nota"])
    inf["etapas"]["capas"] = {"lotes": len(lev), "cotas": len(cotas), "coordenadas": len(coord), "lados": conf_lados}
    por_revisar = [l["nodo"] for l in inf["lotes"] if l["revisar"]]
    inf["por_revisar"] = por_revisar
    inf["requiere_revision"] = inf["requiere_revision"] or bool(por_revisar)
    registrar(f"Capas: {len(lev)} lotes, {len(coord)} vértices; lados {conf_lados}; lotes por revisar: {len(por_revisar)}")
    _vista_general(t, base, lotes_info, inf)

    # 6. Atlas y ZIP
    if t.hacer_atlas:
        nodos = [l["nodo"] for l in inf["lotes"]]
        if cfg.get("salida", "propia") == "qgis":
            _atlas(t, cfg, rutas, inf, nodos, registrar)
        else:
            _salida_propia(t, base, cfg, lev, coord, inf, nodos, registrar)
    return _cerrar(t, inf)


def _salida_propia(t, base, cfg, lev, coord, inf, nodos, registrar):
    """PDF por predio dibujado por el propio motor (sin QGIS)."""
    from . import plano_pdf
    previa = t.vista_previa
    salida = os.path.join(t.carpeta, "00.VISTA_PREVIA" if previa else "02.PDF")
    if os.path.isdir(salida):
        for f in os.listdir(salida):
            os.remove(os.path.join(salida, f))
    # None = los datos del firmante por defecto de la configuración; "" = sin firma
    firma = cfg.get("firma") if t.firma is None else t.firma
    firmante = cfg.get("firmante") if t.firmante is None else t.firmante
    try:
        archivos = plano_pdf.exportar(lev, coord, base, cfg, salida, t.fecha or datetime.date.today().strftime("%d/%m/%Y"),
                                      firma or None, firmante or None,
                                      nodos=[str(n) for n in nodos[:t.nodos_previa]] if previa else None, vista_previa=previa)
    except Exception as e:
        inf["avisos"].append(f"Falló la salida cartográfica: {e}")
        inf["requiere_revision"] = True
        inf["etapas"]["atlas"] = {"ok": False}
        return
    inf["etapas"]["atlas"] = {"ok": True, "archivos": len(archivos), "vista_previa": previa, "motor": "propio"}
    registrar(f"Planos: {len(archivos)} {'vistas previas' if previa else 'PDF'} (render propio)")
    if previa:
        inf["archivos"]["vista_previa"] = archivos
        return
    inf["archivos"]["pdf"] = archivos
    ruta_zip = os.path.join(t.carpeta, f"{t.nombre}.zip")
    with zipfile.ZipFile(ruta_zip, "w", zipfile.ZIP_DEFLATED) as z:
        # 1. PDFs oficiales en 02.PDF
        for a in archivos:
            z.write(a, arcname=os.path.join("02.PDF", os.path.basename(a)))
        # 2. Shapefiles de entrega catastral en 01.SHAPEFILE
        carp_shp = os.path.join(t.carpeta, "01.SHAPEFILE")
        if os.path.isdir(carp_shp):
            for arch in os.listdir(carp_shp):
                z.write(os.path.join(carp_shp, arch), arcname=os.path.join("01.SHAPEFILE", arch))
        # 3. Vista general de control
        vista = os.path.join(t.carpeta, "vista_general.png")
        if os.path.isfile(vista):
            z.write(vista, arcname="vista_general.png")
    inf["archivos"]["zip"] = ruta_zip


def _vista_general(t, base, lotes_info, inf):
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except Exception:
        return
    u = shapely.union_all([i["geom"] for i in lotes_info])
    minx, miny, maxx, maxy = u.buffer(max(12.0, 0.15 * max(u.bounds[2] - u.bounds[0], u.bounds[3] - u.bounds[1]))).bounds
    fig, ax = plt.subplots(figsize=(12, 12 * (maxy - miny) / max(maxx - minx, 1)), dpi=110)
    terr = base.terrenos.cx[minx:maxx, miny:maxy]
    terr.boundary.plot(ax=ax, color="#9a9a9a", linewidth=0.7)
    for r in terr.itertuples():
        c = r.geometry.representative_point()
        if minx < c.x < maxx and miny < c.y < maxy:
            ax.text(c.x, c.y, r.codigo[17:21], fontsize=6, color="#888", ha="center")
    vias = base.vias.cx[minx:maxx, miny:maxy]
    vias.plot(ax=ax, color="#e0457b", linewidth=0.9)
    revisar = {l["nodo"] for l in inf["lotes"] if l.get("revisar")}
    for i in lotes_info:
        g = i["geom"]
        col = "#d9480f" if i["NODO"] in revisar else "#1b5e20"
        ax.fill(*g.exterior.xy, facecolor=col, alpha=0.18, edgecolor=col, linewidth=1.4)
        c = g.representative_point()
        ax.text(c.x, c.y, str(i["NODO"]), fontsize=8, color=col, ha="center", va="center", weight="bold")
    ax.set_xlim(minx, maxx)
    ax.set_ylim(miny, maxy)
    ax.set_aspect("equal")
    ax.set_xticks([])
    ax.set_yticks([])
    ax.set_title("Ubicación de los lotes sobre la base catastral (naranja = revisar)", fontsize=10)
    ruta = os.path.join(t.carpeta, "vista_general.png")
    fig.savefig(ruta, bbox_inches="tight")
    plt.close(fig)
    inf["archivos"]["vista_general"] = ruta


def python_qgis(cfg) -> str:
    ruta = cfg.get("python_qgis") or os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs", "OSGeo4W", "bin",
                                                  "python-qgis-ltr.bat")
    if not os.path.isfile(ruta):
        raise FileNotFoundError("No se encontró el Python de QGIS (python-qgis-ltr.bat). Indique 'python_qgis' en la configuración.")
    return ruta


def _atlas(t: Trabajo, cfg, rutas, inf, nodos, registrar):
    previa = t.vista_previa
    salida = os.path.join(t.carpeta, "00.VISTA_PREVIA" if previa else "02.PDF")
    if os.path.isdir(salida):
        for f in os.listdir(salida):
            os.remove(os.path.join(salida, f))
    job = {
        "proyecto": cfg["proyecto_qgis"], "plantilla": cfg.get("plantilla"), "layout": cfg.get("layout", "SALIDAS_GRAFICAS"),
        "capas": rutas, "capas_mapa": cfg.get("capas_mapa", CAPAS_MAPA), "salida": salida,
        "fecha": t.fecha or datetime.date.today().strftime("%d/%m/%Y"), "firma": t.firma, "firmante": t.firmante,
        "nodos": nodos[:t.nodos_previa] if previa else None, "vista_previa": previa,
        "resultado": os.path.join(t.carpeta, "atlas_resultado.json"),
    }
    ruta_job = os.path.join(t.carpeta, "atlas_trabajo.json")
    with open(ruta_job, "w", encoding="utf-8") as f:
        json.dump(job, f, ensure_ascii=False, indent=1)
    script = os.path.join(RAIZ, "motor", "atlas_qgis.py")
    proc = subprocess.run(["cmd", "/c", python_qgis(cfg), script, ruta_job], capture_output=True, text=True, timeout=3600)
    try:
        with open(job["resultado"], encoding="utf-8") as f:
            res = json.load(f)
    except OSError:
        res = {"ok": False, "error": (proc.stderr or proc.stdout or "sin salida")[-800:]}
    if not res.get("ok"):
        inf["avisos"].append("Falló la salida cartográfica: " + str(res.get("error")))
        inf["requiere_revision"] = True
        inf["etapas"]["atlas"] = {"ok": False}
        return
    archivos = res["archivos"]
    inf["etapas"]["atlas"] = {"ok": True, "archivos": len(archivos), "vista_previa": previa}
    registrar(f"Atlas: {len(archivos)} {'vistas previas' if previa else 'PDF'}")
    if previa:
        inf["archivos"]["vista_previa"] = archivos
        return
    inf["archivos"]["pdf"] = archivos
    ruta_zip = os.path.join(t.carpeta, f"{t.nombre}.zip")
    with zipfile.ZipFile(ruta_zip, "w", zipfile.ZIP_DEFLATED) as z:
        for a in archivos:
            z.write(a, arcname=os.path.join("02.PDF", os.path.basename(a)))
        carp_shp = os.path.join(t.carpeta, "01.SHAPEFILE")
        if os.path.isdir(carp_shp):
            for arch in os.listdir(carp_shp):
                z.write(os.path.join(carp_shp, arch), arcname=os.path.join("01.SHAPEFILE", arch))
        vista = os.path.join(t.carpeta, "vista_general.png")
        if os.path.isfile(vista):
            z.write(vista, arcname="vista_general.png")
    inf["archivos"]["zip"] = ruta_zip


def _cerrar(t: Trabajo, inf):
    ruta = os.path.join(t.carpeta, "informe.json")
    with open(ruta, "w", encoding="utf-8") as f:
        json.dump(inf, f, ensure_ascii=False, indent=1, default=str)
    inf["archivos"]["informe"] = ruta
    try:
        ruta_excel = os.path.join(t.carpeta, "INFORME_REVISION.xlsx")
        escribir_informe_excel(inf, ruta_excel)
        inf["archivos"]["informe_excel"] = ruta_excel
        ruta_zip = inf.get("archivos", {}).get("zip")
        if ruta_zip and os.path.isfile(ruta_zip) and os.path.isfile(ruta_excel):
            with zipfile.ZipFile(ruta_zip, "a") as z:
                z.write(ruta_excel, arcname="INFORME_REVISION.xlsx")
    except Exception as e:      # el informe en Excel es accesorio
        inf["avisos"].append(f"No se pudo escribir INFORME_REVISION.xlsx: {e}")
    return inf


def escribir_informe_excel(inf, ruta):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    col = {"ALTA": "C6EFCE", "MEDIA": "FFEB9C", "BAJA": "FFC7CE"}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LOTES"
    ws.append(["NODO", "REVISAR", "CODIGO", "ORIGEN", "CONFIANZA", "NOTA", "% DENTRO DEL PREDIO", "AREA m2", "FRENTE A VIA"])
    for l in inf["lotes"]:
        ws.append([l["nodo"], "SI" if l.get("revisar") else "", l["codigo"], l["codigo_origen"], l["codigo_confianza"],
                   l["codigo_nota"], round(100 * l["calce"]), l["area"], "SI" if l.get("frente_a_via") else "NO"])
        ws.cell(ws.max_row, 5).fill = PatternFill("solid", fgColor=col.get(l["codigo_confianza"], "FFFFFF"))
        ws.cell(ws.max_row, 3).number_format = "@"
    ws2 = wb.create_sheet("COLINDANTES")
    ws2.append(["NODO", "TRAMO", "DISTANCIA", "LINDERO", "CARDINAL", "COLINDANTE", "TIPO", "CONFIANZA", "% DEL LADO",
                "SEPARACION m", "OTROS COLINDANTES EN EL LADO"])
    for l in inf["lotes"]:
        for s in l.get("lados", []):
            ws2.append([l["nodo"], s["tramo"], s["distancia"], s["lindero"], s["cardinal"], s["colindante"], s["tipo"],
                        s["confianza"], round(100 * s["fraccion"]), s["separacion"],
                        "; ".join(f"{t} ({round(100 * v)}%)" for t, v in s["otros"])])
            ws2.cell(ws2.max_row, 8).fill = PatternFill("solid", fgColor=col.get(s["confianza"], "FFFFFF"))
    ws3 = wb.create_sheet("AVISOS")
    for a in inf["avisos"] or ["Sin avisos."]:
        ws3.append([a])
    ws3.append([])
    ws3.append(["ETAPAS"])
    for k, v in inf["etapas"].items():
        ws3.append([k, json.dumps(v, ensure_ascii=False, default=str)])
    for hoja in (ws, ws2):
        for c in hoja[1]:
            c.font = Font(bold=True)
        hoja.freeze_panes = "A2"
        for letra, ancho in zip("ABCDEFGHIJK", (8, 9, 33, 18, 12, 44, 14, 12, 12, 13, 50)):
            hoja.column_dimensions[letra].width = ancho
    ws3.column_dimensions["A"].width, ws3.column_dimensions["B"].width = 30, 110
    wb.save(ruta)
